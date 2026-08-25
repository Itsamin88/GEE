"""Document parsing: PDF, DOCX, XLSX, CSV, archives, and the failure paths."""

from __future__ import annotations

import io
import zipfile

import pytest

from dcr.extract.dispatch import extract
from dcr.extract.html import parse_date_string, parse_html
from dcr.extract.spreadsheet import extract_csv, extract_xlsx
from dcr.net.mime import sniff
from fixtures.content import make_docx, make_pdf, make_png, make_xlsx, make_zip


@pytest.fixture(scope="module")
def config(request):
    import yaml
    from pathlib import Path

    root = Path(__file__).resolve().parents[1]
    return yaml.safe_load((root / "config" / "config.yaml").read_text(encoding="utf-8"))


# -- MIME sniffing ---------------------------------------------------------
def test_extension_is_never_trusted(config):
    pdf = make_pdf(["Rapport annuel 2019", "Nous cultivons 4 hectares."])
    result = extract(pdf, declared_mime="text/html", filename="page.html", config=config)
    assert result.extension == "pdf"
    assert result.parser_status == "parsed"
    assert "hectares" in result.text


def test_sniff_identifies_ooxml_inside_zip():
    data = make_xlsx({"S": [["a", "b"], [1, 2]]})
    mime, extension = sniff(data, declared="application/octet-stream", filename="x.bin")
    assert extension == "xlsx"


# -- PDF -------------------------------------------------------------------
def test_pdf_text_and_metadata(config):
    pdf = make_pdf(["Rapport annuel 2019", "La foret-jardin couvre 1,8 hectare."],
                   title="Rapport annuel 2019")
    result = extract(pdf, filename="r.pdf", config=config)
    assert result.page_count == 1
    assert "1,8 hectare" in result.text
    assert result.text_status in ("extracted", "ocr_used")


def test_corrupt_pdf_is_recorded_not_raised(config):
    result = extract(b"%PDF-1.4\ntruncated", filename="broken.pdf", config=config)
    assert result.parser_status == "corrupt"
    assert result.text_status == "failed"
    assert result.detail          # the reason is recorded, not swallowed


def test_empty_file_is_handled(config):
    result = extract(b"", filename="empty.pdf", config=config)
    assert result.parser_status in ("corrupt", "unsupported_format", "parsed")


# -- Office ----------------------------------------------------------------
def test_docx_paragraphs_headings_and_tables(config):
    data = make_docx(["Nous avons plante 3000 arbres en 2016."], headings=["Histoire"])
    result = extract(data, filename="plan.docx", config=config)
    assert result.parser_status == "parsed"
    assert "Histoire" in result.headings
    assert "3000 arbres" in result.text


# -- Spreadsheets ----------------------------------------------------------
def test_xlsx_keeps_sheet_and_cell_coordinates():
    data = make_xlsx({"Plantations": [["annee", "surface_ha"], [2016, 1.8], [2019, 0.5]]})
    result = extract_xlsx(data)
    table = result.tables[0]
    assert table.sheet_name == "Plantations"
    assert table.header == ["annee", "surface_ha"]
    assert table.cell_range == "A1:B3"
    assert table.cell_reference(1, 1) == "Plantations!B2"


def test_hidden_sheets_are_inspected_and_flagged():
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "Public"
    workbook.active.append(["a", "b"])
    workbook.active.append([1, 2])
    hidden = workbook.create_sheet("Interne")
    hidden.append(["note", "budget"])
    hidden.append(["plan", 12000])
    hidden.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = extract_xlsx(buffer.getvalue())
    assert "Interne" in result.hidden_sheets
    assert any(t.hidden for t in result.tables)


def test_csv_detects_semicolon_delimiter():
    result = extract_csv("annee;surface\n2016;1,8\n".encode("utf-8"), filename="i.csv")
    assert result.tables[0].header == ["annee", "surface"]
    assert result.tables[0].rows[1] == ["2016", "1,8"]


def test_csv_undecodable_is_reported():
    result = extract_csv(b"\xff\xfe\x00\x00garbage", filename="x.csv")
    assert result.parser_status in ("parsed", "corrupt")


# -- Archives --------------------------------------------------------------
def test_zip_members_are_extracted_for_separate_parsing(config):
    data = make_zip({"note.txt": b"Projet 2017: creation d'une mare."})
    result = extract(data, filename="d.zip", config=config)
    assert result.parser_status == "parsed"
    assert result.contained_files
    assert result.contained_files[0][0] == "note.txt"


def test_zip_path_traversal_is_neutralised(config):
    data = make_zip({"../../etc/passwd": b"x", "ok.txt": b"y"})
    result = extract(data, filename="t.zip", config=config)
    names = [n for n, _ in result.contained_files]
    assert all(".." not in n and "/" not in n for n in names)


def test_decompression_bomb_is_refused(config):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("big.txt", b"0" * 20_000_000)
    result = extract(buffer.getvalue(), filename="bomb.zip", config=config)
    assert result.parser_status == "too_large"
    assert "bomb" in result.detail or "ratio" in result.detail


# -- HTML ------------------------------------------------------------------
def test_html_extracts_links_images_feeds_and_platform():
    html = """<html lang="fr"><head><title>T</title>
    <meta property="article:published_time" content="2017-04-12">
    <link rel="alternate" type="application/rss+xml" href="/feed">
    <script src="/wp-content/x.js"></script></head><body>
    <nav><a href="/about">A</a></nav>
    <figure><img src="/i/plan.jpg" alt="plan de masse"><figcaption>Plan 2016</figcaption></figure>
    <a href="/d/r.pdf">Rapport</a>
    <footer><a href="https://facebook.com/x">FB</a></footer></body></html>"""
    page = parse_html(html, "https://x.org/p")
    assert page.html_lang == "fr"
    assert page.published_date == "2017-04-12"
    assert page.platform_engine == "wordpress"
    assert page.feeds == ["https://x.org/feed"]
    assert page.document_links[0][0] == "https://x.org/d/r.pdf"
    assert page.images[0].caption == "Plan 2016"
    assert page.social_links == ["https://facebook.com/x"]


def test_malformed_html_does_not_raise():
    page = parse_html("<html><body><div><p>unclosed", "https://x.org")
    assert "unclosed" in page.text


@pytest.mark.parametrize("raw, expected", [
    ("2019-03-01", "2019-03-01"),
    ("Mon, 02 Mar 2015 10:00:00 +0000", "2015-03-02"),
    ("12 avril 2017", "2017-04-12"),
    ("March 3, 2011", "2011-03-03"),
    ("sinds 1998", "1998"),
    ("", None),
    ("no date here", None),
])
def test_date_parsing(raw, expected):
    assert parse_date_string(raw) == expected


def test_image_bytes_are_passed_through(config):
    result = extract(make_png(400, 300), filename="p.png", config=config)
    assert result.image_status == "extracted"
    assert result.images[0].extension == "png"
